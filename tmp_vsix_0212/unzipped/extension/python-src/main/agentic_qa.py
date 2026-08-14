import argparse
import asyncio
import csv
import hashlib
import json
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
from urllib.parse import urljoin, urlparse

from playwright.async_api import Error, Page, async_playwright

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document
except ImportError:
    Document = None

CODE_ROOT = Path(__file__).resolve().parent.parent
WORKSPACE_ROOT = Path(os.environ.get("AGENTIC_QA_WORKSPACE_ROOT", str(CODE_ROOT))).resolve()
if str(CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(CODE_ROOT))
from dom import dom as dom_helper

DOCS_DIR = WORKSPACE_ROOT / "docsnew"
DOM_DIR = WORKSPACE_ROOT / "dom" / "dom_elements"
PLAYWRIGHT_DIR = WORKSPACE_ROOT / "playwright-script"
PLAYWRIGHT_POM_DIR = WORKSPACE_ROOT / "playwrightfolder"
PLAYWRIGHT_PAGES_DIR = PLAYWRIGHT_POM_DIR / "pages"
PLAYWRIGHT_TESTS_DIR = PLAYWRIGHT_POM_DIR / "tests"
PLAYWRIGHT_DATA_DIR = PLAYWRIGHT_POM_DIR / "data"
SCREENSHOT_DIR = WORKSPACE_ROOT / "screenshots"
FAILURE_SCREENSHOT_DIR = WORKSPACE_ROOT / "failusecases"
USECASES_DIR = WORKSPACE_ROOT / "usecases"
REPORTS_DIR = WORKSPACE_ROOT / "reports"
RESULTS_DIR = WORKSPACE_ROOT / "results"
CONFIG_DIR = WORKSPACE_ROOT / "configuration"
CONFIG_FILE = CONFIG_DIR / "playwright_config.json"

RESULTS_CSV = RESULTS_DIR / "results.csv"
RESULTS_TXT = RESULTS_DIR / "results.txt"
RESULTS_MD = RESULTS_DIR / "results.md"
RESULTS_JSON = RESULTS_DIR / "results.json"
URL_REGEX = re.compile(r"https?://[\w\-\.\/:?&=#%]+")
SUPPORTED_DOC_SUFFIXES = [".md", ".txt", ".pdf", ".docx"]
DATA_ITEM_REGEX = re.compile(r"^(username|password|email)\s*:\s*(?:\*\*([^*]+)\*\*|(.+))", re.IGNORECASE)
DEFAULT_CLICK_TIMEOUT_SECONDS = 20
DEFAULT_PAGE_DISCOVERY_LIMIT = 10
DEFAULT_USECASE_TOTAL = 50
MAX_USECASE_RETRIES = 1
SCENARIO_STORE_FILE = USECASES_DIR / "scenario_registry.json"
AUTH_LOGIN_KEYWORDS = ("login", "log in", "sign in", "signin", "authentication", "auth")
AUTH_LOGOUT_KEYWORDS = ("logout", "log out", "sign out", "signout")
AUTH_SUCCESS_KEYWORDS = ("dashboard", "welcome", "logged in", "my account", "profile", "secure area", "successfully logged")
AUTH_ERROR_KEYWORDS = ("invalid", "incorrect", "required", "error", "failed", "warning", "try again", "mismatch")
AUTH_FIELD_KEYWORDS = ("user", "username", "email", "password", "login", "signin", "sign in")
SECTION_HEADER_PREFIXES = (
    "scenario",
    "changes required",
    "from the current flow",
    "rules",
    "notes",
    "outputs",
    "purpose",
    "how it works",
)

for path in [
    RESULTS_DIR,
    DOM_DIR,
    PLAYWRIGHT_DIR,
    PLAYWRIGHT_POM_DIR,
    PLAYWRIGHT_PAGES_DIR,
    PLAYWRIGHT_TESTS_DIR,
    PLAYWRIGHT_DATA_DIR,
    SCREENSHOT_DIR,
    FAILURE_SCREENSHOT_DIR,
    USECASES_DIR,
    REPORTS_DIR,
    CONFIG_DIR,
]:
    path.mkdir(parents=True, exist_ok=True)


def slugify(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip().lower())
    return cleaned.strip("_") or "scenario"


def read_pdf_file(path: Path) -> str:
    if PyPDF2 is None:
        raise RuntimeError("PDF support requires PyPDF2. Install it with 'pip install PyPDF2'.")
    reader = PyPDF2.PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_docx_file(path: Path) -> str:
    if Document is None:
        raise RuntimeError("DOCX support requires python-docx. Install it with 'pip install python-docx'.")
    document = Document(str(path))
    parts: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cell for cell in cells if cell))

    return "\n".join(parts)


def read_doc_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return normalize_document_text(read_pdf_file(path))
    if suffix == ".docx":
        return normalize_document_text(read_docx_file(path))
    return normalize_document_text(path.read_text(encoding="utf-8"))


def read_first_doc_file() -> Optional[Path]:
    docs: List[Path] = []
    for suffix in SUPPORTED_DOC_SUFFIXES:
        docs.extend(sorted(DOCS_DIR.glob(f"*{suffix}")))
    return docs[0] if docs else None


def save_scenario_text(text: str, file_name: str = "pasted_scenario.md") -> Path:
    if not text.strip():
        raise ValueError("Scenario text is empty.")
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    scenario_path = DOCS_DIR / file_name
    scenario_path.write_text(text.strip() + "\n", encoding="utf-8")
    return scenario_path


def prompt_user_for_scenario_text() -> str:
    print("Enter the scenario text for testing.")
    print("Start with the target URL on the first line.")
    print("Finish by typing a single line containing only EOF.")
    lines: List[str] = []
    while True:
        try:
            line = input()
        except EOFError:
            break
        if line.strip() == "EOF":
            break
        lines.append(line)
    return "\n".join(lines).strip()


def prompt_user_for_file_path() -> str:
    return input("Enter the path to the .pdf, .md, .txt or .docx instruction file: ").strip()


def resolve_doc_path(doc_path: Optional[str], scenario_text: Optional[str]) -> Optional[Path]:
    if scenario_text:
        if "\\n" in scenario_text:
            scenario_text = scenario_text.replace("\\n", "\n")
        return save_scenario_text(scenario_text)

    if doc_path:
        provided_path = Path(doc_path)
        if not provided_path.is_absolute():
            provided_path = WORKSPACE_ROOT / provided_path
        if not provided_path.exists():
            raise FileNotFoundError(f"Instruction file not found: {provided_path}")
        if provided_path.suffix.lower() not in SUPPORTED_DOC_SUFFIXES:
            raise ValueError(f"Instruction file must be one of: {', '.join(SUPPORTED_DOC_SUFFIXES)}")
        return provided_path

    existing = read_first_doc_file()
    if existing:
        use_existing = input(f"Found existing instruction file '{existing.name}'. Use it? [Y/n]: ").strip().lower()
        if use_existing in ["", "y", "yes"]:
            return existing

    choice = input("No instruction file selected. Type 'file' to provide a path or 'paste' to enter scenario text: ").strip().lower()
    if choice == "file":
        file_path = prompt_user_for_file_path()
        return resolve_doc_path(file_path, None)
    if choice == "paste":
        return save_scenario_text(prompt_user_for_scenario_text())

    print("No valid input selected.")
    return None


def extract_url(text: str) -> Optional[str]:
    first_line = text.strip().splitlines()[0] if text.strip() else ""
    if URL_REGEX.match(first_line):
        return first_line.strip()
    found = URL_REGEX.search(text)
    return found.group(0) if found else None


def normalize_document_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = normalized.replace("\u00a0", " ").replace("\ufeff", "").replace("\u200b", "")
    return normalized.strip() + ("\n" if normalized.strip() else "")


def clean_instruction_line(line: str) -> str:
    cleaned = re.sub(r"^[\-\*\u2022]+\s*", "", line.strip())
    cleaned = re.sub(r"^\d+\s*[\.\)]\s*", "", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip(" -")


def is_instruction_noise(line: str) -> bool:
    lowered = line.lower().strip()
    if not lowered:
        return True
    if lowered.startswith("#"):
        return True
    if lowered.startswith("note:"):
        return True
    if lowered.endswith(":") and not URL_REGEX.search(lowered):
        return True
    if lowered in SECTION_HEADER_PREFIXES:
        return True
    if re.fullmatch(r"[\W_]+", lowered):
        return True
    return False


def is_usecase_generation_directive(line: str) -> bool:
    return bool(
        re.search(
            r"(?:draft|generate|create|prepare|write|build)\s+\d+\s+(?:(?:high\s*level|low\s*level|positive|negative|happy\s*path|edge|exploratory|\w+)\s+){0,5}(?:test\s*)?(?:use\s*cases|cases|test\s*cases)",
            line,
            re.IGNORECASE,
        )
    )


def is_usecase_count_directive(line: str) -> bool:
    return bool(
        re.search(
            r"(?:must\s+not\s+exceed|should\s+not\s+exceed|do\s+not\s+exceed|not\s+exceed(?:ing)?|no\s+more\s+than|up\s+to|maximum\s+of|max(?:imum)?)\s*(?:more\s+than\s*)?\d+\b",
            line,
            re.IGNORECASE,
        )
    )


def parse_instructions(text: str) -> List[str]:
    lines = [clean_instruction_line(line) for line in text.splitlines() if line.strip()]
    instructions = []
    for line in lines:
        if not line:
            continue
        if URL_REGEX.match(line):
            continue
        if DATA_ITEM_REGEX.match(line):
            continue
        if is_instruction_noise(line):
            continue
        if is_usecase_generation_directive(line):
            continue
        if is_usecase_count_directive(line):
            continue
        instructions.append(line)
    return instructions


def extract_metadata(text: str) -> dict:
    metadata = {}
    for line in text.splitlines():
        match = DATA_ITEM_REGEX.match(line.strip())
        if match:
            key = match.group(1).lower()
            metadata[key] = (match.group(2) or match.group(3) or "").strip()
    return metadata


def requested_usecase_count(text: str) -> Optional[int]:
    normalized_text = " ".join(text.split())
    for pattern in [
        r"(?:draft|generate|create|prepare|write|build)\s*(\d+)\s+(?:(?:high\s*level|low\s*level|positive|negative|happy\s*path|edge|exploratory|\w+)\s+){0,5}(?:test\s*)?(?:use\s*cases|cases|test\s*cases)",
        r"(?:must\s+not\s+exceed|should\s+not\s+exceed|do\s+not\s+exceed|not\s+exceed(?:ing)?|no\s+more\s+than|up\s+to|maximum\s+of|max(?:imum)?)\s*(?:more\s+than\s*)?(\d+)\s*(?:test\s*)?(?:use\s*cases|cases|test\s*cases)?",
        r"more than\s*(\d+)\s*(?:use\s*cases|cases)",
        r"at least\s*(\d+)\s*(?:use\s*cases|cases)",
        r"(?:exactly|only|just)\s*(\d+)\s*(?:test\s*)?(?:use\s*cases|cases|test\s*cases)",
        r"(?:need|want|run|execute)\s*(\d+)\s*(?:test\s*)?(?:use\s*cases|cases|test\s*cases)",
        r"(\d+)\s*[:\-]\s*(\d+)\s*(?:use\s*cases|cases)",
        r"(\d+)\s*(?:test\s*)?(?:use\s*cases|cases|test\s*cases)\s*(?:per\s*page)?",
    ]:
        match = re.search(pattern, normalized_text, re.IGNORECASE)
        if not match:
            continue
        numbers = [int(group) for group in match.groups() if group]
        if len(numbers) == 2:
            return max(max(numbers), 3)
        if numbers:
            return max(numbers[0], 3)
    return None


def has_login_credentials(credentials: dict) -> bool:
    return bool(credentials.get("password") and (credentials.get("username") or credentials.get("email")))


def login_identifier_key(credentials: dict) -> str:
    return "username" if credentials.get("username") else "email"


def login_identifier_value(credentials: dict) -> str:
    return credentials.get("username") or credentials.get("email") or ""


def is_login_intent(instructions: List[str], metadata: dict) -> bool:
    lower_text = " ".join(instructions).lower()
    return has_login_credentials(metadata) or any(
        term in lower_text for term in ["login", "sign in", "signin", "credentials", "password", "username", "email"]
    )


def is_form_intent(instructions: List[str]) -> bool:
    lower_text = " ".join(instructions).lower()
    return any(term in lower_text for term in ["fill", "enter", "select", "submit", "search", "input"])


def contains_any_keyword(value: str, keywords: Tuple[str, ...]) -> bool:
    lowered = value.lower()
    return any(keyword in lowered for keyword in keywords)


def normalize_instruction(instruction: str) -> str:
    return instruction.lower().strip()


def scenario_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def load_registry() -> dict:
    if not SCENARIO_STORE_FILE.exists():
        return {}
    try:
        return json.loads(SCENARIO_STORE_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def save_registry(registry: dict) -> None:
    SCENARIO_STORE_FILE.write_text(json.dumps(registry, indent=2), encoding="utf-8")


def build_scenario_id(doc_path: Path, text: str) -> str:
    return f"{slugify(doc_path.stem)}_{scenario_hash(text)[:10]}"


def usecase_bundle_paths(scenario_id: str) -> Dict[str, Path]:
    return {
        "json": USECASES_DIR / f"{scenario_id}_usecases.json",
        "md": USECASES_DIR / f"{scenario_id}_usecases.md",
        "txt": USECASES_DIR / f"{scenario_id}_usecases.txt",
        "csv": USECASES_DIR / f"{scenario_id}_usecases.csv",
    }


def site_understanding_paths(scenario_id: str) -> Dict[str, Path]:
    return {
        "json": USECASES_DIR / f"{scenario_id}_site_understanding.json",
        "md": USECASES_DIR / f"{scenario_id}_site_understanding.md",
        "txt": USECASES_DIR / f"{scenario_id}_site_understanding.txt",
    }


def infer_page_purpose(page_data: dict) -> str:
    title = page_data.get("title", "")
    url = page_data.get("url", "")
    headings = " ".join(page_data.get("headings", []))
    link_text = " ".join((link.get("text") or "") for link in page_data.get("links", []))
    button_text = " ".join((button.get("text") or "") for button in page_data.get("buttons", []))
    source = " ".join([title, url, headings, link_text, button_text]).lower()

    if page_has_login_signal(page_data):
        return "authentication"
    if "contact" in source or "support" in source:
        return "contact"
    if "search" in source:
        return "search"
    if page_data.get("inputs"):
        return "form"
    if "blog" in source or "article" in source:
        return "content"
    if "course" in source or "pricing" in source or "shop" in source:
        return "catalog"
    return "navigation"


def page_has_login_signal(page_data: dict) -> bool:
    title = page_data.get("title", "")
    url = page_data.get("url", "")
    headings = " ".join(page_data.get("headings", []))
    text_preview = page_data.get("text_preview", "")
    inputs = page_data.get("inputs", [])
    buttons = page_data.get("buttons", [])
    source = " ".join([title, url, headings, text_preview]).lower()
    if contains_any_keyword(source, AUTH_LOGIN_KEYWORDS):
        return True
    if any((item.get("type") or "").lower() == "password" for item in inputs):
        return True
    for item in inputs:
        candidate = " ".join([
            item.get("name", ""),
            item.get("id", ""),
            item.get("placeholder", ""),
            item.get("type", ""),
        ]).lower()
        if contains_any_keyword(candidate, AUTH_FIELD_KEYWORDS):
            return True
    for button in buttons:
        if contains_any_keyword(button.get("text", ""), AUTH_LOGIN_KEYWORDS):
            return True
    return False


def page_has_logout_signal(page_data: dict) -> bool:
    title = page_data.get("title", "")
    url = page_data.get("url", "")
    text_preview = page_data.get("text_preview", "")
    source = " ".join([title, url, text_preview]).lower()
    if contains_any_keyword(source, AUTH_LOGOUT_KEYWORDS):
        return True
    for item in page_data.get("links", []) + page_data.get("buttons", []):
        text = item.get("text") or item.get("href") or ""
        if contains_any_keyword(text, AUTH_LOGOUT_KEYWORDS):
            return True
    return False


def find_logout_target(page_data: dict) -> Optional[str]:
    for item in page_data.get("links", []) + page_data.get("buttons", []):
        text = (item.get("text") or "").strip()
        href = (item.get("href") or "").strip()
        if text and contains_any_keyword(text, AUTH_LOGOUT_KEYWORDS):
            return text
        if href and contains_any_keyword(href, AUTH_LOGOUT_KEYWORDS):
            return href
    return None


def infer_document_focus_areas(instructions: List[str], metadata: dict) -> List[str]:
    focus_areas: List[str] = []
    lower_text = " ".join(instructions).lower()
    if is_login_intent(instructions, metadata):
        focus_areas.append("authentication")
    if is_form_intent(instructions):
        focus_areas.append("form validation")
    if "search" in lower_text:
        focus_areas.append("search")
    if "link" in lower_text or "navigate" in lower_text or "menu" in lower_text:
        focus_areas.append("navigation")
    if "logout" in lower_text or "sign out" in lower_text:
        focus_areas.append("logout")
    return focus_areas or ["general navigation"]


def build_site_understanding(
    url: str,
    instructions: List[str],
    metadata: dict,
    discovered_pages: List[dict],
    requested_count: Optional[int],
) -> dict:
    pages = discovered_pages or [{"url": url, "title": "Primary page", "links": [], "buttons": [], "inputs": []}]
    auth_pages = [page for page in pages if page_has_login_signal(page)]
    logout_pages = [page for page in pages if page_has_logout_signal(page)]
    primary_auth_page = auth_pages[0] if auth_pages else pages[0]
    logout_target = None
    for page in auth_pages + logout_pages + pages:
        logout_target = find_logout_target(page)
        if logout_target:
            break

    page_summaries = []
    for page in pages:
        purpose = infer_page_purpose(page)
        page_summaries.append({
            "title": page.get("title") or page.get("url") or "Untitled page",
            "url": page.get("url") or url,
            "purpose": purpose,
            "link_count": len(page.get("links", [])),
            "button_count": len(page.get("buttons", [])),
            "input_count": len(page.get("inputs", [])),
            "headings": page.get("headings", [])[:5],
            "auth_related": page_has_login_signal(page) or page_has_logout_signal(page),
            "text_preview": page.get("text_preview", ""),
        })

    auth_detected = bool(auth_pages) or is_login_intent(instructions, metadata)
    coverage_hints: List[str] = []
    if auth_detected:
        coverage_hints.append("Include positive login coverage with valid credentials.")
        coverage_hints.append("Include logout coverage as part of the high-level auth flow.")
        coverage_hints.append("Include invalid and incomplete credential checks.")
    if any(page.get("input_count", 0) > 0 for page in page_summaries):
        coverage_hints.append("Validate visible form behavior and error feedback.")
    if any(page.get("link_count", 0) > 0 for page in page_summaries):
        coverage_hints.append("Cover major navigation paths across discovered pages.")

    summary = {
        "primary_url": url,
        "requested_usecase_count": requested_count or DEFAULT_USECASE_TOTAL,
        "page_count": len(pages),
        "auth_detected": auth_detected,
        "login_page_url": primary_auth_page.get("url") or url,
        "login_page_title": primary_auth_page.get("title") or "Primary page",
        "logout_target": logout_target or "logout",
        "document_focus_areas": infer_document_focus_areas(instructions, metadata),
        "coverage_hints": coverage_hints,
        "pages": page_summaries,
        "summary_text": (
            f"Discovered {len(pages)} page(s). "
            f"Primary auth page: {(primary_auth_page.get('title') or primary_auth_page.get('url') or url)}. "
            f"Focus areas: {', '.join(infer_document_focus_areas(instructions, metadata))}."
        ),
    }
    return summary


def write_site_understanding_files(scenario_id: str, site_understanding: dict) -> Dict[str, Path]:
    paths = site_understanding_paths(scenario_id)
    lines = [
        "# Site Understanding",
        f"Primary URL: {site_understanding['primary_url']}",
        f"Requested Use Cases: {site_understanding['requested_usecase_count']}",
        f"Discovered Pages: {site_understanding['page_count']}",
        f"Authentication Detected: {'Yes' if site_understanding['auth_detected'] else 'No'}",
        f"Primary Auth Page: {site_understanding['login_page_title']} ({site_understanding['login_page_url']})",
        f"Logout Target Hint: {site_understanding['logout_target']}",
        "",
        f"Summary: {site_understanding['summary_text']}",
        "",
        "Focus Areas:",
    ]
    for area in site_understanding.get("document_focus_areas", []):
        lines.append(f"- {area}")
    lines.append("")
    lines.append("Coverage Hints:")
    for hint in site_understanding.get("coverage_hints", []):
        lines.append(f"- {hint}")
    lines.append("")
    lines.append("Discovered Pages:")
    for page in site_understanding.get("pages", []):
        lines.append(
            f"- {page['title']} | {page['purpose']} | {page['url']} | "
            f"links={page['link_count']} buttons={page['button_count']} inputs={page['input_count']}"
        )
        if page.get("headings"):
            lines.append(f"  headings: {', '.join(page['headings'])}")

    content = "\n".join(lines)
    paths["json"].write_text(json.dumps(site_understanding, indent=2), encoding="utf-8")
    paths["md"].write_text(content, encoding="utf-8")
    paths["txt"].write_text(content, encoding="utf-8")
    return paths


def append_usecase(
    usecases: List[dict],
    title: str,
    description: str,
    case_type: str,
    instructions: List[str],
    page_url: str,
    category: Optional[str] = None,
) -> None:
    if any(existing["title"] == title and existing.get("page_url", "") == page_url for existing in usecases):
        return
    payload = {
        "title": title,
        "description": description,
        "type": case_type,
        "instructions": instructions,
        "page_url": page_url,
    }
    if category:
        payload["category"] = category
    usecases.append(payload)


def build_baseline_usecases(
    url: str,
    instructions: List[str],
    metadata: dict,
    site_understanding: Optional[dict] = None,
) -> List[dict]:
    lower_text = " ".join(instructions).lower()
    usecases: List[dict] = []
    auth_detected = (site_understanding or {}).get("auth_detected", False) or is_login_intent(instructions, metadata)
    auth_page_url = (site_understanding or {}).get("login_page_url", url)
    logout_target = (site_understanding or {}).get("logout_target", "logout")

    if auth_detected:
        has_valid_credentials = has_login_credentials(metadata)
        valid_identifier_key = login_identifier_key(metadata) if has_valid_credentials else "username"
        valid_identifier_value = login_identifier_value(metadata) if has_valid_credentials else ""
        invalid_identifier_value = "invalid_email@example.com" if valid_identifier_key == "email" else "invalid_user"
        login_steps: List[str] = []
        logout_steps: List[str] = []
        if has_valid_credentials:
            login_steps.extend([
                f"fill {valid_identifier_key} with {valid_identifier_value}",
                f"fill password with {metadata.get('password', '')}",
            ])
            logout_steps.extend(login_steps)
        else:
            login_steps.append("verify auth form is visible")
            logout_steps.append("verify auth form is visible")
        login_steps.extend([
            "click login",
            "verify authenticated state",
            "capture DOM after login",
        ])
        logout_steps.extend([
            "click login",
            "verify authenticated state",
            "capture DOM after login",
            f"click {logout_target}",
            "verify logout succeeds",
            "capture DOM after logout",
        ])

        append_usecase(
            usecases,
            "Happy path: positive login",
            "Validate that valid credentials allow the user to log in successfully.",
            "happy",
            login_steps,
            auth_page_url,
            "authentication",
        )
        append_usecase(
            usecases,
            "Happy path: positive login and logout",
            "Validate the complete authentication round trip from login through logout.",
            "happy",
            logout_steps,
            auth_page_url,
            "authentication",
        )
        append_usecase(
            usecases,
            "Negative path: invalid login is rejected",
            "Validate expected error handling when invalid credentials are submitted.",
            "negative",
            [
                f"fill {valid_identifier_key} with {invalid_identifier_value}",
                "fill password with invalid_pass",
                "click login",
                "verify error message appears",
                "capture DOM after invalid login",
            ],
            auth_page_url,
            "authentication",
        )
        append_usecase(
            usecases,
            "Negative path: missing login data is validated",
            "Validate required-field behavior for the authentication form.",
            "negative",
            [
                "submit required inputs incorrectly",
                "verify validation or error feedback",
                "capture DOM after invalid input",
            ],
            auth_page_url,
            "authentication",
        )
        append_usecase(
            usecases,
            "Edge case: authentication controls remain visible",
            "Check that the login page exposes the expected controls before interaction.",
            "edge",
            [
                "verify page loads",
                "verify auth form is visible",
                "verify links are visible",
                "capture DOM after auth form inspection",
            ],
            auth_page_url,
            "authentication",
        )
        return usecases

    if is_form_intent(instructions):
        happy_steps = [
            instr for instr in instructions
            if normalize_instruction(instr).startswith(("fill", "enter", "click", "select", "submit", "search"))
        ] or ["verify page loads", "capture DOM"]
        negative_steps = [
            "submit required inputs incorrectly",
            "verify validation or error feedback",
            "capture DOM after invalid input",
        ]
    else:
        happy_steps = ["verify page loads", "capture DOM"]
        negative_steps = ["verify page loads", "capture DOM"]

    edge_steps = ["verify navigation is stable", "capture DOM after each major step"]
    if "link" in lower_text or "responsive" in lower_text:
        edge_steps.append("verify links are visible")
    if "search" in lower_text:
        edge_steps.append("enter search query with test")

    append_usecase(
        usecases,
        "Happy path: basic flow",
        f"Validate the happy path on {url} using the provided instructions.",
        "happy",
        happy_steps,
        url,
    )
    append_usecase(
        usecases,
        "Negative path: invalid or missing input",
        "Validate expected failures when input is incorrect or required steps are skipped.",
        "negative",
        negative_steps,
        url,
    )
    append_usecase(
        usecases,
        "Edge case: alternate navigation or UI behavior",
        "Verify edge cases and alternate navigation from the same starting point.",
        "edge",
        edge_steps,
        url,
    )
    return usecases


def build_seed_usecases(
    url: str,
    instructions: List[str],
    metadata: dict,
    requested_count: Optional[int] = None,
    site_understanding: Optional[dict] = None,
    expand_to_requested_count: bool = False,
) -> List[dict]:
    usecases = build_baseline_usecases(url, instructions, metadata, site_understanding)

    target_count = max(requested_count or len(usecases), len(usecases))
    if expand_to_requested_count and target_count > len(usecases):
        usecases.extend(
            build_exploratory_cases(
                {"url": url, "title": "Primary page", "links": [], "buttons": [], "inputs": [], "headings": [], "text_preview": ""},
                target_count - len(usecases),
                metadata,
                instructions,
            )
        )
    return usecases


def infer_action_candidates(page_data: dict, primary_url: str) -> List[Tuple[str, str]]:
    actions: List[Tuple[str, str]] = []
    links = page_data.get("links", [])
    buttons = page_data.get("buttons", [])
    inputs = page_data.get("inputs", [])

    for link in links[:20]:
        label = link.get("text") or link.get("href") or "link"
        href = link.get("href") or primary_url
        actions.append(("click", f"{label}||{href}"))

    for button in buttons[:15]:
        label = button.get("text") or button.get("selector") or "button"
        actions.append(("button", label))

    for input_item in inputs[:15]:
        label = input_item.get("name") or input_item.get("placeholder") or input_item.get("type") or "field"
        actions.append(("fill", label))

    return actions


def build_exploratory_cases(page_data: dict, target_count: int, metadata: dict, instructions: List[str]) -> List[dict]:
    page_title = page_data.get("title") or page_data.get("url") or "Page"
    page_url = page_data.get("url") or ""
    actions = infer_action_candidates(page_data, page_url)
    exploratory_cases: List[dict] = []
    login_hint = page_has_login_signal(page_data) or is_login_intent(instructions, metadata)

    for index in range(target_count):
        action = actions[index % len(actions)] if actions else ("observe", "page")
        steps = ["verify page loads", "capture DOM before actions"]
        action_type, action_target = action
        title_suffix = ""

        if action_type == "click":
            label, href = action_target.split("||", 1)
            steps.append(f"click {href}")
            steps.append("verify page loads")
            steps.append("capture DOM after click")
            title_suffix = f"click link {label[:30]}"
        elif action_type == "button":
            steps.append(f"click {action_target}")
            steps.append("capture DOM after click")
            title_suffix = f"click action {action_target[:30]}"
        elif action_type == "fill":
            value = "edge_case_input" if index % 3 == 0 else ("invalid_input" if index % 3 == 1 else "happy_path_input")
            steps.append(f"fill {action_target} with {value}")
            steps.append("capture DOM after input")
            title_suffix = f"fill field {action_target[:30]}"
        else:
            steps.append("verify navigation is stable")
            steps.append("capture DOM after each major step")
            title_suffix = "observe page"

        if login_hint and index % 4 == 0:
            steps.append("verify auth form is visible")

        if index % 5 == 0:
            steps.append("verify links are visible")
        if index % 7 == 0:
            steps.append("submit required inputs incorrectly")

        exploratory_cases.append({
            "title": f"{page_title}: exploratory case {index + 1}",
            "description": f"Auto-generated exploratory case for {page_title} to {title_suffix}.",
            "type": "exploratory",
            "instructions": steps,
            "page_url": page_url,
            "category": title_suffix,
        })

    return exploratory_cases


def build_final_usecases(
    url: str,
    instructions: List[str],
    metadata: dict,
    discovered_pages: List[dict],
    requested_count: Optional[int],
    site_understanding: Optional[dict] = None,
) -> List[dict]:
    final_usecases = build_seed_usecases(
        url,
        instructions,
        metadata,
        requested_count=None,
        site_understanding=site_understanding,
        expand_to_requested_count=False,
    )
    target_count = max(requested_count, len(final_usecases)) if requested_count is not None else DEFAULT_USECASE_TOTAL
    pages = discovered_pages or [{"url": url, "title": "Primary page", "links": [], "buttons": [], "inputs": [], "headings": [], "text_preview": ""}]
    remaining = target_count - len(final_usecases)
    if remaining <= 0:
        return final_usecases

    base_count = remaining // len(pages)
    extra = remaining % len(pages)
    for index, page in enumerate(pages):
        page_target = base_count + (1 if index < extra else 0)
        if page_target <= 0:
            continue
        final_usecases.extend(build_exploratory_cases(page, page_target, metadata, instructions))

    return final_usecases


def write_usecases_files(scenario_id: str, source_name: str, url: str, usecases: List[dict]) -> Dict[str, Path]:
    paths = usecase_bundle_paths(scenario_id)
    lines = [f"# Use cases generated from {source_name}", f"URL: {url}", ""]
    for uc in usecases:
        lines.append(f"## {uc['title']}")
        lines.append(f"Type: {uc['type']}")
        lines.append(f"Description: {uc['description']}")
        if uc.get("page_url"):
            lines.append(f"Page URL: {uc['page_url']}")
        lines.append("Instructions:")
        for inst in uc["instructions"]:
            lines.append(f"- {inst}")
        lines.append("")

    paths["md"].write_text("\n".join(lines), encoding="utf-8")
    paths["txt"].write_text("\n".join(lines), encoding="utf-8")
    with paths["csv"].open("w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(["title", "type", "description", "page_url", "instructions"])
        for uc in usecases:
            writer.writerow([uc["title"], uc["type"], uc["description"], uc.get("page_url", ""), " | ".join(uc["instructions"])])

    paths["json"].write_text(json.dumps(usecases, indent=2), encoding="utf-8")
    return paths


def load_saved_usecases(scenario_id: str) -> Optional[List[dict]]:
    path = usecase_bundle_paths(scenario_id)["json"]
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


def prompt_for_approval(usecases: List[dict], scenario_id: str, auto_approve: bool, site_understanding: Optional[dict] = None) -> bool:
    if auto_approve:
        print(f"Auto-approve enabled for scenario {scenario_id}.")
        return True

    if site_understanding:
        print("I have understood what this site is about and stored the discovered DOM inventory.")
        print(site_understanding.get("summary_text", ""))
        print(
            "Discovered pages: "
            + ", ".join(page["title"] for page in site_understanding.get("pages", [])[:5])
        )
    print(f"Generated {len(usecases)} use cases for scenario '{scenario_id}'.")
    print("Review the files under /usecases and approve before browser execution.")
    approval = input("Approve generated use cases and continue? [Y/N]: ").strip().lower()
    return approval in {"y", "yes"}


def update_registry_entry(scenario_id: str, doc_path: Path, text: str, approved: bool, usecase_count: int) -> None:
    registry = load_registry()
    registry[scenario_id] = {
        "scenario_id": scenario_id,
        "source_file": str(doc_path),
        "hash": scenario_hash(text),
        "approved": approved,
        "usecase_count": usecase_count,
        "updated_at": datetime.now().isoformat(),
    }
    save_registry(registry)


async def timed_click(page: Page, target: str, click_timeout_seconds: int) -> bool:
    if not target:
        return False
    selectors = []
    if target.startswith("http://") or target.startswith("https://"):
        selectors.extend([f"a[href='{target}']", f"a[href=\"{target}\"]", target, f"text={target}"])
    elif target.startswith("xpath=") or target.startswith("//") or target.startswith("/"):
        selectors.append(target if target.startswith("xpath=") else f"xpath={target}")
    else:
        selectors.extend([
            target,
            f"text={target}",
            f"button:has-text(\"{target}\")",
            f"a:has-text(\"{target}\")",
            f"[aria-label*='{target}']",
            f"input[value*='{target}']",
        ])

    for selector in selectors:
        try:
            locator = page.locator(selector)
            if await locator.count() > 0:
                await asyncio.wait_for(locator.first.click(), timeout=click_timeout_seconds)
                return True
        except Exception:
            continue
        try:
            await asyncio.wait_for(page.click(selector), timeout=click_timeout_seconds)
            return True
        except Exception:
            continue
    return False


async def safe_fill(page: Page, target: str, value: str) -> bool:
    if not target:
        return False
    if target.startswith("xpath=") or target.startswith("//") or target.startswith("/"):
        selector = target if target.startswith("xpath=") else f"xpath={target}"
        try:
            await page.fill(selector, value)
            return True
        except Exception:
            return False

    selectors = []
    if target.startswith("#") or target.startswith(".") or target.startswith("css=") or "[" in target or "=" in target:
        selectors.append(target)
    selectors += [
        f"input[name*='{target}']",
        f"input[id*='{target}']",
        f"input[placeholder*='{target}']",
        f"input[aria-label*='{target}']",
        f"textarea[name*='{target}']",
        f"textarea[id*='{target}']",
        f"textarea[placeholder*='{target}']",
    ]

    for selector in selectors:
        try:
            await page.fill(selector, value)
            return True
        except Exception:
            continue
    return False


async def fill_first_matching_field(page: Page, selectors: List[str], value: str) -> bool:
    for selector in selectors:
        try:
            locator = page.locator(selector)
            if await locator.count() > 0:
                await locator.first.fill(value)
                return True
        except Exception:
            continue
    return False


async def read_body_text(page: Page) -> str:
    try:
        return await page.locator("body").inner_text(timeout=5000)
    except Exception:
        return ""


async def is_auth_form_visible(page: Page) -> bool:
    try:
        password_count = await page.locator("input[type='password']").count()
        identifier_count = await page.locator(
            "input[type='text'], input[type='email'], input[name*='user'], input[id*='user'], input[name*='email'], input[id*='email']"
        ).count()
        return password_count > 0 and identifier_count > 0
    except Exception:
        return False


async def page_has_error_feedback(page: Page) -> bool:
    try:
        alert_count = await page.locator(
            "[role='alert'], .error, .alert, .validation-error, .message-error, .woocommerce-error"
        ).count()
        if alert_count > 0:
            return True
    except Exception:
        pass
    body_text = (await read_body_text(page)).lower()
    return contains_any_keyword(body_text, AUTH_ERROR_KEYWORDS)


async def page_looks_authenticated(page: Page) -> bool:
    body_text = (await read_body_text(page)).lower()
    current_url = page.url.lower()
    auth_form_visible = await is_auth_form_visible(page)
    if contains_any_keyword(body_text, AUTH_LOGOUT_KEYWORDS):
        return True
    if contains_any_keyword(body_text, AUTH_SUCCESS_KEYWORDS):
        return True
    if not auth_form_visible and not contains_any_keyword(current_url, AUTH_LOGIN_KEYWORDS):
        return True
    return False


async def page_looks_logged_out(page: Page) -> bool:
    body_text = (await read_body_text(page)).lower()
    if await is_auth_form_visible(page):
        return True
    if contains_any_keyword(body_text, AUTH_LOGIN_KEYWORDS):
        return True
    return False


async def attempt_form_submission(page: Page, click_timeout_seconds: int) -> bool:
    for target in ["login", "sign in", "submit", "continue"]:
        if await timed_click(page, target, click_timeout_seconds):
            return True
    return False


def extract_login_credentials_from_text(text: str) -> dict:
    credentials: dict = {}
    normalized = re.sub(r"[ \t]+", " ", text)
    lines = [re.sub(r"\s+", " ", line).strip() for line in normalized.splitlines() if line.strip()]
    candidate_text = "\n".join(
        line for line in lines
        if any(keyword in line.lower() for keyword in ["username", "email", "password", "credential", "login"])
    ) or normalized

    email_match = re.search(
        r"(?:email|e-mail)\s*(?::|is|=)?\s*([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})",
        candidate_text,
        re.IGNORECASE,
    )
    if email_match:
        credentials["email"] = email_match.group(1).strip()

    username_match = re.search(
        r"(?:username|user\s*name|user id|userid)\s*(?::|is|=)?\s*([A-Z0-9._-]+)",
        candidate_text,
        re.IGNORECASE,
    )
    if username_match:
        candidate = username_match.group(1).strip()
        if candidate.lower() not in {"username", "user", "email", "login"}:
            credentials["username"] = candidate

    password_match = re.search(
        r"(?:password|pass\s*word)\s*(?::|is|=)?\s*([^\s,;]+)",
        candidate_text,
        re.IGNORECASE,
    )
    if password_match:
        candidate = password_match.group(1).strip()
        if candidate.lower() not in {"password", "pass", "login"}:
            credentials["password"] = candidate

    return credentials


async def discover_site_login_credentials(page: Page) -> dict:
    try:
        body_text = await page.locator("body").inner_text(timeout=5000)
    except Exception:
        return {}
    return extract_login_credentials_from_text(body_text)


async def resolve_login_credentials(page: Page, metadata: dict, instructions: List[str]) -> dict:
    resolved = dict(metadata)
    if has_login_credentials(resolved):
        return resolved
    if not is_login_intent(instructions, metadata):
        return resolved

    discovered = await discover_site_login_credentials(page)
    for key in ["username", "email", "password"]:
        if discovered.get(key) and not resolved.get(key):
            resolved[key] = discovered[key]
    return resolved


async def fill_login_form(page: Page, credentials: dict) -> bool:
    if not has_login_credentials(credentials):
        return False

    identifier_value = login_identifier_value(credentials)
    identifier_key = login_identifier_key(credentials)
    identifier_selectors = [
        "input[name*=username]",
        "input[id*=username]",
        "input[name*=user]",
        "input[id*=user]",
        "input[name*=email]",
        "input[id*=email]",
        "input[placeholder*=user]",
        "input[placeholder*=email]",
        "input[type=email]",
        "input[type=text]",
    ]
    password_selectors = [
        "input[type=password]",
        "input[name*=password]",
        "input[id*=password]",
        "input[placeholder*=password]",
    ]

    filled_identifier = await fill_first_matching_field(page, identifier_selectors, identifier_value)
    if not filled_identifier:
        filled_identifier = await safe_fill(page, identifier_key, identifier_value)
    if not filled_identifier and identifier_key != "username":
        filled_identifier = await safe_fill(page, "username", identifier_value)
    if not filled_identifier and identifier_key != "email":
        filled_identifier = await safe_fill(page, "email", identifier_value)

    filled_password = await fill_first_matching_field(page, password_selectors, credentials["password"])
    if not filled_password:
        filled_password = await safe_fill(page, "password", credentials["password"])

    return filled_identifier or filled_password


class POMPage:
    def __init__(self, page: Page, dom_dir: Path):
        self.page = page
        self.dom_dir = dom_dir

    async def goto(self, url: str, timeout_ms: int) -> None:
        await self.page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
        await self.page.wait_for_load_state("networkidle", timeout=timeout_ms)
        await self.capture_dom("home")

    async def capture_dom(self, name: str) -> str:
        content = await self.page.content()
        file_path = self.dom_dir / f"{datetime.now():%Y%m%d_%H%M%S_%f}_{slugify(name)}.html"
        file_path.write_text(content, encoding="utf-8")
        dom_helper.save_dom_snapshot(name, content)
        return str(file_path.relative_to(WORKSPACE_ROOT))

    async def screenshot(self, name: str, failed: bool = False) -> str:
        SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)
        FAILURE_SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)
        filename = SCREENSHOT_DIR / f"{datetime.now():%Y%m%d_%H%M%S_%f}_{slugify(name)}{'_FAIL' if failed else '_PASS'}.png"
        await self.page.screenshot(path=filename, full_page=True)
        if failed:
            failure_copy = FAILURE_SCREENSHOT_DIR / filename.name
            shutil.copyfile(filename, failure_copy)
        return str(filename.relative_to(WORKSPACE_ROOT))


async def collect_page_inventory(page: Page) -> dict:
    return await page.evaluate(
        """() => {
            const links = Array.from(document.querySelectorAll('a'))
                .map((link, index) => ({
                    text: (link.innerText || link.textContent || '').trim().replace(/\\s+/g, ' ').slice(0, 120),
                    href: link.href || '',
                    selector: `a:nth-of-type(${index + 1})`
                }))
                .filter(link => link.text || link.href);
            const buttons = Array.from(document.querySelectorAll('button, input[type="submit"], input[type="button"]'))
                .map((button, index) => ({
                    text: (button.innerText || button.value || button.getAttribute('aria-label') || '').trim().replace(/\\s+/g, ' ').slice(0, 120),
                    selector: button.id ? `#${button.id}` : `${button.tagName.toLowerCase()}:nth-of-type(${index + 1})`
                }))
                .filter(button => button.text || button.selector);
            const inputs = Array.from(document.querySelectorAll('input, textarea, select'))
                .map((input, index) => ({
                    name: input.getAttribute('name') || '',
                    id: input.getAttribute('id') || '',
                    placeholder: input.getAttribute('placeholder') || '',
                    type: input.getAttribute('type') || input.tagName.toLowerCase(),
                    selector: input.id ? `#${input.id}` : `${input.tagName.toLowerCase()}:nth-of-type(${index + 1})`
                }));
            const headings = Array.from(document.querySelectorAll('h1, h2, h3'))
                .map((heading) => (heading.innerText || heading.textContent || '').trim().replace(/\\s+/g, ' ').slice(0, 120))
                .filter(Boolean)
                .slice(0, 8);
            const textPreview = ((document.body && document.body.innerText) || '')
                .trim()
                .replace(/\\s+/g, ' ')
                .slice(0, 280);
            return {
                title: document.title || '',
                url: window.location.href,
                links,
                buttons,
                inputs,
                headings,
                text_preview: textPreview
            };
        }"""
    )


async def discover_pages_and_dom(page: Page, start_url: str, timeout_ms: int, max_pages: int) -> List[dict]:
    parsed_start = urlparse(start_url)
    same_host = parsed_start.netloc
    queue: List[str] = [start_url]
    visited: Set[str] = set()
    pages: List[dict] = []

    while queue and len(visited) < max_pages:
        current = queue.pop(0)
        if current in visited:
            continue
        visited.add(current)
        try:
            await page.goto(current, wait_until="domcontentloaded", timeout=timeout_ms)
            await page.wait_for_load_state("networkidle", timeout=timeout_ms)
            content = await page.content()
            dom_helper.save_dom_snapshot(slugify(current), content)
            inventory = await collect_page_inventory(page)
            inventory["discovered_at"] = datetime.now().isoformat()
            pages.append(inventory)

            for link in inventory.get("links", []):
                href = link.get("href") or ""
                if not href:
                    continue
                full_url = urljoin(current, href)
                parsed = urlparse(full_url)
                if parsed.scheme not in {"http", "https"}:
                    continue
                if parsed.netloc != same_host:
                    continue
                if full_url not in visited and full_url not in queue and len(queue) + len(visited) < max_pages * 2:
                    queue.append(full_url)
        except Exception as exc:
            pages.append({
                "url": current,
                "title": "Discovery failed",
                "links": [],
                "buttons": [],
                "inputs": [],
                "error": str(exc),
                "discovered_at": datetime.now().isoformat(),
            })

    DOM_DIR.mkdir(parents=True, exist_ok=True)
    inventory_path = DOM_DIR / "dom_inventory.json"
    inventory_path.write_text(json.dumps(pages, indent=2), encoding="utf-8")
    return pages


def write_playwright_assets(scenario_id: str, url: str, usecases: List[dict]) -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    flat_script_path = PLAYWRIGHT_DIR / f"{scenario_id}_{timestamp}_playwright.py"
    (PLAYWRIGHT_PAGES_DIR / "__init__.py").write_text("", encoding="utf-8")
    (PLAYWRIGHT_TESTS_DIR / "__init__.py").write_text("", encoding="utf-8")
    flat_script = f"""import asyncio
from pathlib import Path
from playwright.async_api import async_playwright

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
USECASES = {json.dumps(usecases, indent=2)}
URL = {url!r}

async def main():
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=False)
        context = await browser.new_context()
        page = await context.new_page()
        for usecase in USECASES:
            await page.goto(usecase.get("page_url") or URL, wait_until="networkidle")
            screenshot_name = f"{{usecase['title']}}_{{usecase['type']}}".replace(" ", "_").lower()
            await page.screenshot(path=str(WORKSPACE_ROOT / "screenshots" / f"{{screenshot_name}}_sample.png"), full_page=True)
        await browser.close()

if __name__ == "__main__":
    asyncio.run(main())
"""
    flat_script_path.write_text(flat_script, encoding="utf-8")

    base_page_path = PLAYWRIGHT_PAGES_DIR / "base_page.py"
    if not base_page_path.exists():
        base_page_path.write_text(
            """from pathlib import Path
from playwright.async_api import Page

class BasePage:
    def __init__(self, page: Page):
        self.page = page

    async def open(self, url: str, timeout_ms: int = 30000):
        await self.page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
        await self.page.wait_for_load_state("networkidle", timeout=timeout_ms)

    async def screenshot(self, path: Path):
        await self.page.screenshot(path=path, full_page=True)
""",
            encoding="utf-8",
        )

    site_page_path = PLAYWRIGHT_PAGES_DIR / f"{scenario_id}_page.py"
    site_page_path.write_text(
        f"""from .base_page import BasePage

class ScenarioPage(BasePage):
    scenario_id = {scenario_id!r}
    default_url = {url!r}
""",
        encoding="utf-8",
    )

    data_path = PLAYWRIGHT_DATA_DIR / f"{scenario_id}_usecases.json"
    data_path.write_text(json.dumps(usecases, indent=2), encoding="utf-8")

    test_path = PLAYWRIGHT_TESTS_DIR / f"test_{scenario_id}.py"
    test_path.write_text(
        f"""import asyncio
import json
from pathlib import Path
from playwright.async_api import async_playwright
from pages.{scenario_id}_page import ScenarioPage

WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
USECASE_FILE = WORKSPACE_ROOT / "playwrightfolder" / "data" / "{scenario_id}_usecases.json"

async def main():
    usecases = json.loads(USECASE_FILE.read_text(encoding="utf-8"))
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=False)
        context = await browser.new_context()
        page = await context.new_page()
        pom = ScenarioPage(page)
        for usecase in usecases:
            await pom.open(usecase.get("page_url") or pom.default_url)
            screenshot_name = f"{{usecase['title']}}_{{usecase['type']}}".replace(" ", "_").lower()
            await pom.screenshot(WORKSPACE_ROOT / "screenshots" / f"{{screenshot_name}}_qa.png")
        await browser.close()

if __name__ == "__main__":
    asyncio.run(main())
""",
        encoding="utf-8",
    )


def build_runtime_report(
    scenario_id: str,
    source_file: Path,
    url: str,
    usecases: List[dict],
    discovered_pages: List[dict],
    results: List[dict],
    site_understanding: Optional[dict] = None,
) -> dict:
    return {
        "scenario_id": scenario_id,
        "source_file": str(source_file),
        "url": url,
        "generated_at": datetime.now().isoformat(),
        "usecase_count": len(usecases),
        "discovered_page_count": len(discovered_pages),
        "summary": summarize_results(results),
        "results": results,
        "discovered_pages": discovered_pages,
        "site_understanding": site_understanding or {},
    }


def write_results(results: List[dict]) -> None:
    summary = summarize_results(results)
    txt_lines = [
        "TEST RESULTS SUMMARY",
        f"PASS: {summary['passed']} | FAIL: {summary['failed']} | TOTAL: {summary['total']}",
        "",
    ]
    md_lines = [
        "# Test Results",
        "",
        f"Generated: {datetime.now().isoformat()}",
        f"PASS: {summary['passed']} | FAIL: {summary['failed']} | TOTAL: {summary['total']}",
        "",
    ]

    with RESULTS_CSV.open("w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(["title", "type", "result", "status_detail", "page_url", "error", "screenshot", "duration_seconds"])
        for result in results:
            writer.writerow([
                result["title"],
                result["type"],
                result["result"],
                result["status_detail"],
                result.get("page_url", ""),
                result["error"],
                result.get("screenshot", ""),
                result.get("duration_seconds", ""),
            ])
            txt_lines.append(
                f"{result['result']} | {result['title']} | {result['type']} | {result['status_detail']} | {result['error']}"
            )
            md_lines.append(f"## {result['title']}")
            md_lines.append(f"- Type: {result['type']}")
            md_lines.append(f"- Result: {result['result']}")
            md_lines.append(f"- Status Detail: {result['status_detail']}")
            md_lines.append(f"- Page: {result.get('page_url', '')}")
            md_lines.append(f"- Duration Seconds: {result.get('duration_seconds', '')}")
            if result["error"]:
                md_lines.append(f"- Error: {result['error']}")
            md_lines.append(f"- Screenshot: {result.get('screenshot', '')}")
            if result.get("console_errors"):
                md_lines.append(f"- Console Errors: {len(result['console_errors'])}")
            if result.get("network_errors"):
                md_lines.append(f"- Network Errors: {len(result['network_errors'])}")
            md_lines.append("")

    RESULTS_TXT.write_text("\n".join(txt_lines), encoding="utf-8")
    RESULTS_MD.write_text("\n".join(md_lines), encoding="utf-8")
    RESULTS_JSON.write_text(json.dumps(results, indent=2), encoding="utf-8")


def write_reports(report_data: dict) -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = REPORTS_DIR / f"report_{timestamp}.json"
    html_path = REPORTS_DIR / f"extent_report_{timestamp}.html"

    json_path.write_text(json.dumps(report_data, indent=2), encoding="utf-8")

    rows = []
    for result in report_data["results"]:
        result_class = "pass" if result["result"] == "PASS" else "fail"
        rows.append(
            "<tr>"
            f"<td>{result['title']}</td>"
            f"<td>{result['type']}</td>"
            f"<td class='{result_class}'>{result['result']}</td>"
            f"<td>{result['status_detail']}</td>"
            f"<td>{result.get('page_url', '')}</td>"
            f"<td>{result.get('duration_seconds', '')}</td>"
            f"<td>{result.get('error', '')}</td>"
            "</tr>"
        )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <title>Agentic-QA Extent Report</title>
  <style>
    :root {{
      --bg: #f4efe6;
      --card: #fffaf2;
      --ink: #1e2430;
      --accent: #c65d2f;
      --pass: #2f855a;
      --warn: #b7791f;
      --fail: #c53030;
      --border: #e6d6bf;
    }}
    body {{ font-family: Georgia, 'Trebuchet MS', serif; background: linear-gradient(135deg, #f7f1e3, #efe4cf); color: var(--ink); margin: 0; padding: 24px; }}
    .card {{ background: var(--card); border: 1px solid var(--border); border-radius: 20px; padding: 24px; box-shadow: 0 18px 40px rgba(64, 52, 35, 0.08); }}
    h1, h2 {{ margin-top: 0; }}
    .meta {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(220px, 1fr)); gap: 12px; margin-bottom: 24px; }}
    .pill {{ padding: 12px 14px; border-radius: 14px; background: #fff; border: 1px solid var(--border); }}
    table {{ width: 100%; border-collapse: collapse; background: #fff; border-radius: 16px; overflow: hidden; }}
    th, td {{ padding: 12px; border-bottom: 1px solid var(--border); text-align: left; vertical-align: top; }}
    th {{ background: #f8ebd6; }}
    .pass {{ color: var(--pass); font-weight: 700; }}
    .fail {{ color: var(--fail); font-weight: 700; }}
  </style>
</head>
<body>
  <div class="card">
    <h1>Agentic-QA Extent Report</h1>
    <div class="meta">
      <div class="pill"><strong>Scenario:</strong> {report_data['scenario_id']}</div>
      <div class="pill"><strong>Source:</strong> {report_data['source_file']}</div>
      <div class="pill"><strong>URL:</strong> {report_data['url']}</div>
      <div class="pill"><strong>Generated:</strong> {report_data['generated_at']}</div>
      <div class="pill"><strong>Use Cases:</strong> {report_data['usecase_count']}</div>
      <div class="pill"><strong>Discovered Pages:</strong> {report_data['discovered_page_count']}</div>
      <div class="pill"><strong>PASS:</strong> {report_data['summary']['passed']} | <strong>FAIL:</strong> {report_data['summary']['failed']} | <strong>TOTAL:</strong> {report_data['summary']['total']}</div>
    </div>
    <h2>Execution Summary</h2>
    <table>
      <thead>
        <tr>
          <th>Title</th>
          <th>Type</th>
          <th>Result</th>
          <th>Status Detail</th>
          <th>Page URL</th>
          <th>Duration Seconds</th>
          <th>Error</th>
        </tr>
      </thead>
      <tbody>
        {''.join(rows)}
      </tbody>
    </table>
  </div>
</body>
</html>"""
    html_path.write_text(html, encoding="utf-8")


def load_playwright_config() -> dict:
    default = {
        "browser": "chromium",
        "headless": False,
        "timeout": 30000,
        "slow_mo": 250,
        "click_timeout_seconds": DEFAULT_CLICK_TIMEOUT_SECONDS,
        "max_discovery_pages": DEFAULT_PAGE_DISCOVERY_LIMIT,
    }
    if CONFIG_FILE.exists():
        try:
            loaded = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            default.update(loaded)
        except Exception:
            pass
    return default


def clear_run_artifacts() -> None:
    for path in [
        RESULTS_DIR,
        REPORTS_DIR,
        SCREENSHOT_DIR,
        FAILURE_SCREENSHOT_DIR,
        DOM_DIR,
        USECASES_DIR,
        PLAYWRIGHT_DIR,
        PLAYWRIGHT_DATA_DIR,
        PLAYWRIGHT_TESTS_DIR,
        PLAYWRIGHT_PAGES_DIR,
    ]:
        if path.exists():
            for item in path.glob("*"):
                if item.is_dir():
                    shutil.rmtree(item)
                else:
                    item.unlink()
        path.mkdir(parents=True, exist_ok=True)

    for old_file in [WORKSPACE_ROOT / "dom" / "latest_dom.txt", WORKSPACE_ROOT / "dom" / "latest_dom.py"]:
        if old_file.exists():
            old_file.unlink()

    if SCENARIO_STORE_FILE.exists():
        SCENARIO_STORE_FILE.unlink()


def parse_cli_args():
    parser = argparse.ArgumentParser(description="Run the agentic QA automation framework.")
    parser.add_argument("--file", "-f", dest="doc_path", help="Path to a .pdf, .md, .txt or .docx scenario file.")
    parser.add_argument("--scenario", "-s", dest="scenario_text", help="Paste the scenario text directly.")
    parser.add_argument("--approve", action="store_true", help="Automatically approve generated use cases and continue execution.")
    parser.add_argument("--regenerate", action="store_true", help="Retained for CLI compatibility. Fresh generated outputs are written on every run.")
    return parser.parse_args()


def parse_fill_instruction(instruction: str) -> Optional[Tuple[str, str]]:
    match = re.match(r"(?:fill|enter)\s+(.+?)\s+with\s+(.+)", instruction, re.IGNORECASE)
    if not match:
        return None
    return match.group(1).strip(), match.group(2).strip()


async def attach_runtime_listeners(page: Page) -> dict:
    listener_state = {"bucket": None}

    def active_bucket() -> Optional[dict]:
        return listener_state.get("bucket")

    def handle_console(msg) -> None:
        bucket = active_bucket()
        if bucket is not None and msg.type == "error":
            bucket["console"].append({"type": msg.type, "text": msg.text})

    def handle_page_error(exc) -> None:
        bucket = active_bucket()
        if bucket is not None:
            bucket["page_errors"].append(str(exc))

    page.on("console", handle_console)
    page.on("pageerror", handle_page_error)

    async def handle_response(response):
        bucket = active_bucket()
        if bucket is not None and response.status >= 400:
            bucket["network"].append({"url": response.url, "status": response.status})

    page.on("response", lambda response: asyncio.create_task(handle_response(response)))
    return listener_state


def format_result_label(status: str) -> str:
    return "PASS" if status == "passed" else "FAIL"


def summarize_results(results: List[dict]) -> dict:
    passed = sum(1 for result in results if result.get("result") == "PASS")
    failed = sum(1 for result in results if result.get("result") == "FAIL")
    return {"passed": passed, "failed": failed, "total": len(results)}


def decorate_result_record(result: dict) -> dict:
    result["result"] = format_result_label(result["status"])
    result["status_detail"] = result["status"].upper()
    return result


async def execute_instruction(
    pom_page: POMPage,
    instruction: str,
    metadata: dict,
    config: dict,
    result: dict,
) -> None:
    normalized = normalize_instruction(instruction)
    click_timeout_seconds = int(config.get("click_timeout_seconds", DEFAULT_CLICK_TIMEOUT_SECONDS))

    if "click" in normalized:
        target = re.sub(r"(?i)^click\s+", "", instruction).strip()
        started = asyncio.get_running_loop().time()
        clicked = await timed_click(pom_page.page, target, click_timeout_seconds)
        elapsed = asyncio.get_running_loop().time() - started
        if elapsed > click_timeout_seconds:
            raise TimeoutError(f"Click action exceeded {click_timeout_seconds} seconds: {target}")
        if not clicked:
            result["status"] = "warning"
            result["steps"].append(f"Could not click target: {target}")
        else:
            await pom_page.page.wait_for_load_state("networkidle", timeout=config.get("timeout", 30000))
            result["step_timings"].append({"instruction": instruction, "duration_seconds": round(elapsed, 2)})
            await pom_page.capture_dom(f"after_click_{target}")
        return

    if "fill" in normalized or "enter" in normalized:
        parsed = parse_fill_instruction(instruction)
        if not parsed:
            result["status"] = "warning"
            result["steps"].append(f"Skipped unclear fill instruction: {instruction}")
            return
        selector, value = parsed
        if value.lower() == "invalid_input":
            value = "invalid_input_123"
        elif value.lower() == "edge_case_input":
            value = "!@# edge boundary 999"
        elif value.lower() == "happy_path_input":
            value = "happy_path_value"
        filled = await safe_fill(pom_page.page, selector, value)
        if not filled:
            result["status"] = "warning"
            result["steps"].append(f"Could not fill target: {selector}")
        else:
            await pom_page.capture_dom(f"after_fill_{selector}")
        return

    if "submit required inputs incorrectly" in normalized:
        await attempt_form_submission(pom_page.page, click_timeout_seconds)
        await pom_page.capture_dom("invalid_submission_state")
        return

    if "verify auth form is visible" in normalized:
        if not await is_auth_form_visible(pom_page.page):
            raise AssertionError("Authentication form is not visible.")
        await pom_page.capture_dom("auth_form_visible")
        return

    if "verify authenticated state" in normalized:
        if not await page_looks_authenticated(pom_page.page):
            raise AssertionError("Authenticated state was not detected after login.")
        await pom_page.capture_dom("authenticated_state")
        return

    if "verify logout succeeds" in normalized:
        if not await page_looks_logged_out(pom_page.page):
            raise AssertionError("Logout state was not detected.")
        await pom_page.capture_dom("after_logout_verification")
        return

    if "verify error message appears" in normalized or "verify validation or error feedback" in normalized:
        if not await page_has_error_feedback(pom_page.page):
            raise AssertionError("Expected validation or error feedback was not found.")
        await pom_page.capture_dom("validation_feedback")
        return

    if "verify links are visible" in normalized:
        result["link_count"] = await pom_page.page.locator("a").count()
        return

    if "verify page loads" in normalized or "verify navigation is stable" in normalized or "check" in normalized:
        await pom_page.page.wait_for_load_state("networkidle", timeout=config.get("timeout", 30000))
        await pom_page.capture_dom("verification")
        return

    if "capture dom" in normalized:
        await pom_page.capture_dom(normalized.replace(" ", "_"))
        return

    if is_login_intent([instruction], metadata) and has_login_credentials(metadata):
        await fill_login_form(pom_page.page, metadata)
        await pom_page.capture_dom("login_form")


async def run_test_case(page: Page, listener_state: dict, url: str, usecase: dict, metadata: dict, config: dict) -> dict:
    event_bucket = {"console": [], "network": [], "page_errors": []}
    listener_state["bucket"] = event_bucket
    pom_page = POMPage(page, DOM_DIR)

    result = {
        "title": usecase["title"],
        "type": usecase["type"],
        "status": "passed",
        "steps": [],
        "error": "",
        "page_url": usecase.get("page_url") or url,
        "console_errors": [],
        "network_errors": [],
        "page_errors": [],
        "step_timings": [],
        "duration_seconds": 0,
        "screenshot": "",
    }

    started = asyncio.get_running_loop().time()
    try:
        await pom_page.goto(result["page_url"], config.get("timeout", 30000))
        resolved_credentials = await resolve_login_credentials(page, metadata, usecase["instructions"])
        if has_login_credentials(resolved_credentials) and is_login_intent(usecase["instructions"], resolved_credentials):
            await fill_login_form(page, resolved_credentials)
            if not has_login_credentials(metadata):
                result["steps"].append("Detected login credentials from the page")
            result["steps"].append("Filled login credentials")

        for instruction in usecase["instructions"]:
            await execute_instruction(pom_page, instruction, resolved_credentials, config, result)
            result["steps"].append(instruction)

        result["screenshot"] = await pom_page.screenshot(f"{usecase['title']}_{usecase['type']}", failed=False)
    except TimeoutError as exc:
        result["status"] = "failed"
        result["error"] = str(exc)
        result["screenshot"] = await pom_page.screenshot(f"{usecase['title']}_{usecase['type']}", failed=True)
    except Error as exc:
        result["status"] = "failed"
        result["error"] = str(exc)
        result["screenshot"] = await pom_page.screenshot(f"{usecase['title']}_{usecase['type']}", failed=True)
    except Exception as exc:
        result["status"] = "failed"
        result["error"] = str(exc)
        try:
            result["screenshot"] = await pom_page.screenshot(f"{usecase['title']}_{usecase['type']}", failed=True)
        except Exception:
            result["screenshot"] = ""
    finally:
        result["duration_seconds"] = round(asyncio.get_running_loop().time() - started, 2)
        result["console_errors"] = event_bucket["console"]
        result["network_errors"] = event_bucket["network"]
        result["page_errors"] = event_bucket["page_errors"]
        decorate_result_record(result)
        listener_state["bucket"] = None
    return result


async def run_case_with_retry(page: Page, listener_state: dict, url: str, usecase: dict, metadata: dict, config: dict) -> dict:
    timeout_seconds = max(int(config.get("click_timeout_seconds", DEFAULT_CLICK_TIMEOUT_SECONDS)) * 2, 45)
    for attempt in range(1, MAX_USECASE_RETRIES + 2):
        try:
            result = await asyncio.wait_for(
                run_test_case(page, listener_state, url, usecase, metadata, config),
                timeout=timeout_seconds,
            )
            if attempt > 1 and result["status"] != "failed":
                result["steps"].append("Retry attempt succeeded")
            return result
        except asyncio.TimeoutError:
            print(f"Use case '{usecase['title']}' timed out, attempt {attempt}/{MAX_USECASE_RETRIES + 1}.")
            if attempt == MAX_USECASE_RETRIES + 1:
                return decorate_result_record({
                    "title": usecase["title"],
                    "type": usecase["type"],
                    "status": "failed",
                    "steps": [f"Use case timed out after {timeout_seconds} seconds"],
                    "error": f"Timeout after {timeout_seconds} seconds",
                    "page_url": usecase.get("page_url") or url,
                    "console_errors": [],
                    "network_errors": [],
                    "page_errors": [],
                    "step_timings": [],
                    "duration_seconds": timeout_seconds,
                    "screenshot": "",
                })
            print("Retrying use case...")


async def main(
    doc_path: Optional[str] = None,
    scenario_text: Optional[str] = None,
    auto_approve: bool = False,
    regenerate: bool = False,
) -> None:
    try:
        instruction_path = resolve_doc_path(doc_path, scenario_text)
    except (FileNotFoundError, ValueError, RuntimeError) as exc:
        print(str(exc))
        return

    if not instruction_path:
        print("No instruction file found in /docsnew. Add a .md, .txt, .pdf, or .docx file, provide a file path, or paste a scenario.")
        return

    text = read_doc_file(instruction_path)
    url = extract_url(text)
    if not url:
        print("No URL found in the instruction document. Put the target URL on the first line.")
        return

    instructions = parse_instructions(text)
    metadata = extract_metadata(text)
    clear_run_artifacts()

    config = load_playwright_config()
    scenario_id = build_scenario_id(instruction_path, text)
    requested_count = requested_usecase_count(text)
    seed_usecases = build_seed_usecases(url, instructions, metadata, requested_count, expand_to_requested_count=False)
    write_usecases_files(scenario_id, instruction_path.name, url, seed_usecases)
    print(f"Provisional use cases written under /usecases for scenario '{scenario_id}'.")
    usecases = seed_usecases

    print(f"Launching browser headless={config.get('headless', False)} slow_mo={config.get('slow_mo', 250)}")
    playwright = await async_playwright().start()
    browser_launcher = getattr(playwright, config.get("browser", "chromium"))
    browser = await browser_launcher.launch(headless=config.get("headless", False), slow_mo=config.get("slow_mo", 250))
    context = await browser.new_context()

    discovered_pages: List[dict] = []
    site_understanding: Optional[dict] = None
    results: List[dict] = []
    execution_page: Optional[Page] = None
    try:
        discovery_page = await context.new_page()
        discovered_pages = await discover_pages_and_dom(
            discovery_page,
            url,
            timeout_ms=config.get("timeout", 30000),
            max_pages=int(config.get("max_discovery_pages", DEFAULT_PAGE_DISCOVERY_LIMIT)),
        )
        await discovery_page.close()

        site_understanding = build_site_understanding(url, instructions, metadata, discovered_pages, requested_count)
        write_site_understanding_files(scenario_id, site_understanding)

        usecases = build_final_usecases(url, instructions, metadata, discovered_pages, requested_count, site_understanding)
        write_usecases_files(scenario_id, instruction_path.name, url, usecases)
        if not prompt_for_approval(usecases, scenario_id, auto_approve, site_understanding):
            update_registry_entry(scenario_id, instruction_path, text, approved=False, usecase_count=len(usecases))
            print("Use case approval declined. Review the generated files under /usecases and rerun when ready.")
            return
        update_registry_entry(scenario_id, instruction_path, text, approved=True, usecase_count=len(usecases))

        write_playwright_assets(scenario_id, url, usecases)

        write_results(results)
        execution_page = await context.new_page()
        listener_state = await attach_runtime_listeners(execution_page)
        for usecase in usecases:
            result = await run_case_with_retry(execution_page, listener_state, url, usecase, metadata, config)
            results.append(result)
            write_results(results)
            print(f"[{len(results)}/{len(usecases)}] {result['result']} - {result['title']}")
    finally:
        if execution_page is not None and not execution_page.is_closed():
            await execution_page.close()
        await browser.close()
        await playwright.stop()

    write_results(results)
    report_data = build_runtime_report(scenario_id, instruction_path, url, usecases, discovered_pages, results, site_understanding)
    write_reports(report_data)
    print(f"Execution complete. Results written to {RESULTS_DIR}, reports to {REPORTS_DIR}, and Playwright assets to {PLAYWRIGHT_POM_DIR}.")


def main_sync() -> None:
    args = parse_cli_args()
    asyncio.run(main(args.doc_path, args.scenario_text, args.approve, args.regenerate))


if __name__ == "__main__":
    main_sync()
